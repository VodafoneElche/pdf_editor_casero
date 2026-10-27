"""Pasar de PDF a Word (.docx) y de Word a PDF.

- PDF → Word: con pdf2docx, que reconstruye párrafos, tablas e imágenes. El
  .docx se puede subir a Google Drive y abrir con Documentos de Google.
- Word → PDF: si el ordenador tiene Microsoft Word, se usa Word (resultado
  idéntico). Si no, se convierte con mammoth (texto, tablas, listas e
  imágenes) y PyMuPDF; el diseño puede variar un poco.
"""

from __future__ import annotations

import io
import multiprocessing
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf

EXTENSIONES_WORD = (".docx",)


# --- PDF → Word ------------------------------------------------------------------
def pdf_a_word(origen: str, destino: str) -> None:
    """Convierte el PDF `origen` en el .docx `destino` (puede tardar en PDF largos)."""
    import logging

    from pdf2docx import Converter

    logging.getLogger().setLevel(logging.WARNING)  # pdf2docx escribe mucho en el registro
    conversor = Converter(origen)
    try:
        conversor.convert(destino)
    finally:
        conversor.close()


def _proceso_pdf_a_word(origen: str, destino: str, cola) -> None:
    try:
        pdf_a_word(origen, destino)
        cola.put(("ok", destino))
    except Exception as e:  # el mensaje lo enseña la ventana
        cola.put(("error", f"{type(e).__name__}: {e}"))


class ExportacionWord:
    """Exportación en un proceso aparte, para que la ventana no se congele y se pueda cancelar."""

    def __init__(self, datos_pdf: bytes, destino: str):
        self._tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
        self._tmp.write(datos_pdf)
        self._tmp.close()
        contexto = multiprocessing.get_context("spawn")
        self._cola = contexto.Queue()
        self._proceso = contexto.Process(
            target=_proceso_pdf_a_word, args=(self._tmp.name, destino, self._cola), daemon=True
        )
        self._proceso.start()

    def resultado(self):
        """None mientras trabaja; luego ("ok", ruta) o ("error", mensaje)."""
        if not self._cola.empty():
            r = self._cola.get()
            self._limpiar()
            return r
        if not self._proceso.is_alive():
            self._limpiar()
            return ("error", "La conversión se ha detenido sin terminar.")
        return None

    def cancelar(self):
        if self._proceso.is_alive():
            self._proceso.terminate()
        self._limpiar()

    def _limpiar(self):
        try:
            os.remove(self._tmp.name)
        except OSError:
            pass


# --- PDF → hoja de cálculo ---------------------------------------------------------
_NUMERO = re.compile(r"^-?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?\s*(€|%)?$")


def _valor(texto: str):
    """Convierte "1.028,26 €" en el número 1028.26 (formato español); el resto queda como texto."""
    t = " ".join(texto.split())
    if not _NUMERO.match(t):
        return t or None
    simbolo = t[-1] if t[-1] in "€%" else ""
    numero = float(t.rstrip("€% ").replace(".", "").replace(",", "."))
    return (numero / 100 if simbolo == "%" else numero), simbolo


def pdf_a_excel(datos_pdf: bytes, destino: str, cancelado=lambda: False, avance=lambda n, total: None) -> int:
    """Pasa las tablas del PDF a un .xlsx (una hoja por página con tablas). Devuelve cuántas tablas.

    Si no hay ninguna tabla, cada línea de texto va a una fila, para no entregar una hoja vacía.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    doc = pymupdf.open(stream=datos_pdf, filetype="pdf")
    libro = Workbook()
    libro.remove(libro.active)
    negrita = Font(bold=True)
    tablas = 0
    for n, pagina in enumerate(doc):
        if cancelado():
            raise InterruptedError
        avance(n, doc.page_count)
        encontradas = pagina.find_tables().tables
        if not encontradas:
            continue
        hoja = libro.create_sheet(f"Página {n + 1}")
        fila = 1
        for tabla in encontradas:
            tablas += 1
            valores = [
                [_valor(pagina.get_text("text", clip=pymupdf.Rect(c))) if c else None for c in f.cells]
                for f in tabla.rows
            ]
            # Las columnas vacías salen de rayas sueltas del diseño: fuera
            usadas = [j for j in range(tabla.col_count) if any(j < len(f) and f[j] is not None for f in valores)]
            for i, fila_tabla in enumerate(valores):
                for k, j in enumerate(usadas):
                    valor = fila_tabla[j] if j < len(fila_tabla) else None
                    if valor is None:
                        continue
                    destino_celda = hoja.cell(fila, k + 1)
                    if isinstance(valor, tuple):
                        destino_celda.value = valor[0]
                        destino_celda.number_format = {"€": '#,##0.00 "€"', "%": "0.00%"}.get(valor[1], "General")
                    else:
                        destino_celda.value = valor
                    if i == 0:
                        destino_celda.font = negrita
                fila += 1
            fila += 1  # una fila en blanco entre tablas
        for columna in hoja.columns:
            ancho = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
            hoja.column_dimensions[columna[0].column_letter].width = min(max(ancho + 2, 8), 60)
    if tablas == 0:
        for n, pagina in enumerate(doc):
            hoja = libro.create_sheet(f"Página {n + 1}")
            for linea in pagina.get_text("text").splitlines():
                if linea.strip():
                    valor = _valor(linea)
                    hoja.append([valor[0] if isinstance(valor, tuple) else valor])
        if not libro.sheetnames:
            libro.create_sheet("Hoja 1")
    libro.save(destino)
    return tablas


# --- Word → PDF ------------------------------------------------------------------
_SCRIPT_WORD = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
    $doc = $word.Documents.Open($args[0], $false, $true)
    $doc.ExportAsFixedFormat($args[1], 17)
    $doc.Close($false)
} finally {
    $word.Quit()
}
"""


def _con_microsoft_word(origen: str, destino: str) -> bool:
    """Convierte con Word si está instalado (solo Windows). True si ha funcionado."""
    if sys.platform != "win32":
        return False
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
             "-Command", _SCRIPT_WORD, origen, destino],
            capture_output=True, timeout=180, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        return r.returncode == 0 and Path(destino).is_file() and Path(destino).stat().st_size > 0
    except Exception:
        return False


CSS_WORD = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.3; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
p { margin: 0 0 6pt 0; }
table { border-collapse: collapse; margin: 4pt 0; }
td, th { border: 0.5pt solid #999; padding: 2pt 4pt; vertical-align: top; }
img { max-width: 100%; }
"""


def _con_mammoth(origen: str) -> bytes:
    import mammoth

    with open(origen, "rb") as f:
        html = mammoth.convert_to_html(f).value
    historia = pymupdf.Story(html=html or "<p></p>", user_css=CSS_WORD)
    salida = io.BytesIO()
    escritor = pymupdf.DocumentWriter(salida)
    hoja = pymupdf.paper_rect("a4")
    zona = hoja + (56, 56, -56, -56)  # márgenes de 2 cm
    quedan = True
    while quedan:
        dispositivo = escritor.begin_page(hoja)
        quedan, _ = historia.place(zona)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    return salida.getvalue()


def word_a_pdf(origen: str, destino: str) -> str:
    """Convierte el .docx `origen` en el PDF `destino`. Devuelve qué se usó: "Word" o "mammoth"."""
    if _con_microsoft_word(str(Path(origen).resolve()), str(Path(destino).resolve())):
        return "Word"
    Path(destino).write_bytes(_con_mammoth(origen))
    return "mammoth"
