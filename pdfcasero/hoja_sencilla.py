"""PDF → hoja de cálculo «como iLovePDF»: cada dato en su casilla y las imágenes encima.

Cada página va en su pestaña. Las columnas y las filas salen de dónde empiezan los
textos (y dónde acaban los importes, que van alineados a la derecha), con el mismo
ancho y alto que en el PDF, así la hoja se parece a la página. Los fondos de color se
pintan en las casillas, y los logotipos, códigos de barras y gráficas se ponen como
imágenes flotando encima de las casillas, en su sitio.
"""

from __future__ import annotations

import io
import re

import pymupdf

# Lo que Excel no admite dentro de una casilla (los códigos de barras traen de esto)
_PROHIBIDOS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_CIFRAS = re.compile(r"^[-+]?\d[\d.,:]*\s?[€%]?$")
_PUNTOS_A_ANCHO = 96 / 72 / 7  # ancho de columna de Excel: unos 7 píxeles por unidad
_PUNTOS_A_PIXELES = 96 / 72


class Trozo:
    """Un texto suelto de la página: «Nº de factura:», «CI0925062425», «824,10»…"""

    def __init__(self, texto, rect, tamano, negrita, color):
        self.texto, self.rect, self.tamano, self.negrita, self.color = texto, rect, tamano, negrita, color

    @property
    def derecha(self) -> bool:
        """Importes, cantidades, horas y duraciones: van alineados a la derecha."""
        t = self.texto
        if t.isdigit() and len(t) > 4:
            return False  # números de cuenta, teléfonos y códigos: como texto, a la izquierda
        return isinstance(_numero(t), tuple) or bool(_CIFRAS.match(t))


def _numero(texto):
    """Importes y porcentajes como número; los números sin decimales (cuentas, teléfonos,
    códigos) se quedan como texto para que no pierdan ceros ni salgan en notación rara."""
    from .conversion import _valor

    valor = _valor(texto)
    if isinstance(valor, tuple) and ("," not in texto and not valor[1]):
        return texto
    return valor


def _trozos(pagina, zonas_imagen) -> list[Trozo]:
    invisibles = [pymupdf.Rect(t["bbox"]) for t in pagina.get_texttrace() if t.get("type") == 3]
    trozos = []
    for bloque in pagina.get_text("dict")["blocks"]:
        for linea in bloque.get("lines", []):
            if abs(linea["dir"][0] - 1) > 0.01:
                continue  # texto girado: va como imagen
            actual = None
            for s in linea["spans"]:
                texto = _PROHIBIDOS.sub("", s["text"])
                r = pymupdf.Rect(s["bbox"])
                centro = pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
                if not texto.strip():
                    if actual is not None and not actual.texto.endswith(" "):
                        actual.texto += " "  # espacio entre palabras: sigue el mismo texto
                        actual.rect |= r
                    continue
                if any(centro in z for z in zonas_imagen + invisibles):
                    actual = None
                    continue
                negrita = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                # Trozos seguidos casi pegados son el mismo texto; con hueco, otra casilla
                # (en los párrafos justificados el hueco entre palabras llega a ser de más de una letra)
                if actual is not None and r.x0 - actual.rect.x1 < max(3.0, s["size"] * 1.6):
                    hueco = " " if r.x0 - actual.rect.x1 > s["size"] * 0.1 and not actual.texto.endswith(" ") else ""
                    actual.texto += hueco + texto
                    actual.rect |= r
                    continue
                actual = Trozo(texto, r, s["size"], negrita, s["color"])
                trozos.append(actual)
    for t in trozos:
        t.texto = " ".join(t.texto.split())
    return [t for t in trozos if t.texto]


def _agrupar(valores, tolerancia):
    """Junta posiciones casi iguales y devuelve una por grupo."""
    grupos = []
    for v in sorted(valores):
        if grupos and v - grupos[-1][-1] <= tolerancia:
            grupos[-1].append(v)
        else:
            grupos.append([v])
    return [min(g) for g in grupos]


def _columnas(trozos, izquierda, derecha):
    # Bordes: donde empiezan los textos y donde acaban los importes
    bordes = [t.rect.x1 + 1 if t.derecha else t.rect.x0 for t in trozos] + [izquierda, derecha]
    bordes = _agrupar(bordes, 4)
    limpios = [bordes[0]]
    for b in bordes[1:]:
        if b - limpios[-1] >= 6:  # columnas finísimas no: se juntan con la anterior
            limpios.append(b)
    return limpios


def _filas(trozos, arriba, abajo):
    inicios = _agrupar([t.rect.y0 for t in trozos] + [arriba, abajo], 2.5)
    limpios = [inicios[0]]
    for y in inicios[1:]:
        if y - limpios[-1] >= 4:
            limpios.append(y)
    # Debajo de cada renglón, el hueco hasta el siguiente va en una fila vacía aparte: así
    # la fila del texto mide lo que el texto y los fondos de color caen en su sitio
    con_huecos = []
    for k, y in enumerate(limpios):
        con_huecos.append(y)
        if k + 1 == len(limpios):
            break
        fondo = max((t.rect.y1 for t in trozos if y - 0.5 <= t.rect.y0 + 2 < limpios[k + 1]), default=None)
        if fondo is not None and limpios[k + 1] - (fondo + 1.5) >= 4:
            con_huecos.append(fondo + 1.5)
    return con_huecos


def _indice(bordes, x) -> int:
    """La columna (o fila) en la que cae la posición x."""
    i = 0
    while i + 1 < len(bordes) and bordes[i + 1] <= x + 0.5:
        i += 1
    return min(i, len(bordes) - 2) if len(bordes) > 1 else 0


def _hex(color) -> str:
    return "%02X%02X%02X" % tuple(max(0, min(255, round(c * 255))) for c in color)


def _imagenes(pagina, tablas_zonas, fondos):
    from .word_sencillo import _graficas

    zonas = list(_graficas(pagina, tablas_zonas, fondos))
    for img in pagina.get_image_info():
        r = pymupdf.Rect(img["bbox"]) & pagina.rect
        if r.is_empty or r.width < 4 or r.height < 4:
            continue
        if r.width > 0.9 * pagina.rect.width and r.height > 0.9 * pagina.rect.height:
            continue  # imagen de fondo de toda la página (escaneos): el texto va en casillas
        if not any(r in z for z in zonas):
            zonas.append(r)
    return zonas


def pdf_a_hoja(datos_pdf: bytes, destino: str, cancelado=lambda: False, avance=lambda n, total: None) -> int:
    """Escribe el .xlsx con una pestaña por página. Devuelve cuántas casillas con datos hay."""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ImagenExcel
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    from .conversion import preparar_para_word
    from .word_sencillo import _fondos

    datos, _ = preparar_para_word(datos_pdf)  # códigos de barras y texto girado, como imagen
    doc = pymupdf.open(stream=datos, filetype="pdf")
    libro = Workbook()
    libro.remove(libro.active)
    casillas = 0
    for n, pagina in enumerate(doc):
        if cancelado():
            raise InterruptedError
        avance(n, doc.page_count)
        hoja = libro.create_sheet(f"Página {n + 1}")
        hoja.sheet_view.showGridLines = False
        # Al imprimir o pasar a PDF, la página entera a lo ancho de una hoja, como el original
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 1
        hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
        fondos = [(r, c) for r, c in _fondos(pagina) if r.width < 0.95 * pagina.rect.width or r.height < 0.95 * pagina.rect.height]
        try:
            tablas = [pymupdf.Rect(t.bbox) for t in pagina.find_tables(use_layout=False).tables]
        except TypeError:
            tablas = [pymupdf.Rect(t.bbox) for t in pagina.find_tables().tables]
        zonas = _imagenes(pagina, tablas, fondos)
        trozos = _trozos(pagina, zonas)
        if not trozos and not zonas:
            continue
        cosas = [t.rect for t in trozos] + zonas
        izquierda = max(0, min(r.x0 for r in cosas) - 4)
        derecha = max(r.x1 for r in cosas) + 4
        arriba = max(0, min(r.y0 for r in cosas) - 4)
        abajo = max(r.y1 for r in cosas) + 4
        columnas = _columnas(trozos, izquierda, derecha)
        filas = _filas(trozos, arriba, abajo)
        if columnas[-1] < derecha:
            columnas.append(derecha)
        if filas[-1] < abajo:
            filas.append(abajo)
        for j in range(len(columnas) - 1):
            ancho = (columnas[j + 1] - columnas[j]) * _PUNTOS_A_ANCHO
            hoja.column_dimensions[get_column_letter(j + 1)].width = max(0.6, round(ancho, 2))
        for i in range(len(filas) - 1):
            hoja.row_dimensions[i + 1].height = max(3, round(filas[i + 1] - filas[i], 1))
        # Fondos de color: las casillas que quedan dentro se pintan
        for r, color in sorted(fondos, key=lambda f: -f[0].get_area()):
            relleno = PatternFill("solid", fgColor=color)
            for i in range(len(filas) - 1):
                cy = (filas[i] + filas[i + 1]) / 2
                if not r.y0 - 1 <= cy <= r.y1 + 1:
                    continue
                for j in range(len(columnas) - 1):
                    cx = (columnas[j] + columnas[j + 1]) / 2
                    if r.x0 - 1 <= cx <= r.x1 + 1:
                        hoja.cell(i + 1, j + 1).fill = relleno
        # Cada texto en su casilla
        for t in sorted(trozos, key=lambda t: t.derecha):  # los importes al final
            fila = _indice(filas, t.rect.y0 + 2) + 1
            if t.derecha:
                col = _indice(columnas, t.rect.x1 - 1) + 1
                # Los números no se salen de su casilla: si no caben, se juntan las de la izquierda
                desde = col
                ocupadas = {(r.min_row, c) for r in hoja.merged_cells.ranges for c in range(r.min_col, r.max_col + 1)}

                def cabe(c):
                    return columnas[col] - columnas[c - 1] >= t.rect.width * 1.1 + 2

                # sin pasar de donde empieza el número, para no tapar la etiqueta de al lado
                tope = _indice(columnas, t.rect.x0 + 1) + 1
                def libre(c):
                    return c >= 1 and hoja.cell(fila, c).value is None and (fila, c) not in ocupadas

                # primero hasta donde empieza el número; si aun así no cabe, alguna más si está vacía
                while desde > tope and libre(desde - 1):
                    desde -= 1
                while not cabe(desde) and libre(desde - 1):
                    desde -= 1
                if desde < col:
                    hoja.merge_cells(start_row=fila, start_column=desde, end_row=fila, end_column=col)
                    col = desde
            else:
                col = _indice(columnas, t.rect.x0) + 1
            celda = hoja.cell(fila, col)
            valor = _numero(t.texto)
            if celda.value is not None:  # dos textos en la misma casilla: se juntan
                celda.value = f"{celda.value} {t.texto}"
                continue
            if isinstance(valor, tuple):
                celda.value = valor[0]
                decimales = len(t.texto.split(",")[1].rstrip("€% ")) if "," in t.texto else 0
                patron = "#,##0" + ("." + "0" * decimales if decimales else "")
                celda.number_format = {"€": patron + ' "€"', "%": "0.00%"}.get(valor[1], patron)
                celda.alignment = Alignment(horizontal="right", vertical="center")
            else:
                celda.value = valor
                celda.alignment = Alignment(horizontal="right" if t.derecha else None, vertical="center")
            celda.font = Font(name="Arial", size=max(5, round(t.tamano * 2) / 2), bold=t.negrita, color=_hex(_rgb_int(t.color)))
            casillas += 1
        # Imágenes encima de las casillas, en su sitio
        for r in zonas:
            foto = pagina.get_pixmap(clip=r, dpi=150)
            imagen = ImagenExcel(io.BytesIO(foto.tobytes("png")))
            ancho_px, alto_px = r.width * _PUNTOS_A_PIXELES, r.height * _PUNTOS_A_PIXELES
            imagen.width, imagen.height = ancho_px, alto_px
            col = _indice(columnas, r.x0)
            fila = _indice(filas, r.y0)
            desplaza_x = pixels_to_EMU((r.x0 - columnas[col]) * _PUNTOS_A_PIXELES)
            desplaza_y = pixels_to_EMU((r.y0 - filas[fila]) * _PUNTOS_A_PIXELES)
            imagen.anchor = OneCellAnchor(
                _from=AnchorMarker(col=col, colOff=max(0, desplaza_x), row=fila, rowOff=max(0, desplaza_y)),
                ext=XDRPositiveSize2D(pixels_to_EMU(ancho_px), pixels_to_EMU(alto_px)),
            )
            hoja.add_image(imagen)
    if not libro.sheetnames:
        libro.create_sheet("Hoja 1")
    libro.save(destino)
    return casillas


def _rgb_int(color: int):
    return ((color >> 16) & 255) / 255, ((color >> 8) & 255) / 255, (color & 255) / 255
