import pymupdf

from pdfcasero import conversion


def test_pdf_a_word_y_vuelta(tmp_path):
    pdf = tmp_path / "carta.pdf"
    doc = pymupdf.open()
    pagina = doc.new_page()
    pagina.insert_text((72, 80), "Carta para la familia", fontsize=18)
    pagina.insert_htmlbox(pymupdf.Rect(72, 110, 400, 140), "<p>Precio: 25 € al mes</p>")
    doc.save(pdf)

    docx = tmp_path / "carta.docx"
    conversion.pdf_a_word(str(pdf), str(docx))
    assert docx.stat().st_size > 1000

    vuelta = tmp_path / "vuelta.pdf"
    metodo = conversion.word_a_pdf(str(docx), str(vuelta))
    assert metodo in ("Word", "mammoth")
    texto = pymupdf.open(vuelta)[0].get_text()
    assert "Carta para la familia" in texto and "25 €" in texto


def test_exportacion_en_proceso_aparte(tmp_path):
    import time

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 80), "Hola", fontsize=12)
    destino = tmp_path / "hola.docx"
    trabajo = conversion.ExportacionWord(doc.tobytes(), str(destino))
    limite = time.time() + 120
    resultado = None
    while resultado is None and time.time() < limite:
        time.sleep(0.2)
        resultado = trabajo.resultado()
    assert resultado == ("ok", str(destino)), resultado
    assert destino.stat().st_size > 1000


def test_tablas_del_pdf_a_hoja_de_calculo(tmp_path):
    from openpyxl import load_workbook

    from pdfcasero.conversion import _valor, pdf_a_excel

    assert _valor("1.028,26 €") == (1028.26, "€")
    assert _valor("21%") == (0.21, "%")
    assert _valor("Total: 5,00€") == "Total: 5,00€"
    doc = pymupdf.open()
    pagina = doc.new_page()
    filas = [("Concepto", "Importe"), ("Base imponible", "824,10"), ("Total a pagar", "1.028,26 €")]
    for x in (50, 250, 450):  # rejilla de la tabla
        pagina.draw_line((x, 50), (x, 50 + 30 * len(filas)))
    for i in range(len(filas) + 1):
        pagina.draw_line((50, 50 + 30 * i), (450, 50 + 30 * i))
    for i, (a, b) in enumerate(filas):
        pagina.insert_htmlbox(pymupdf.Rect(55, 55 + 30 * i, 245, 78 + 30 * i), a)
        pagina.insert_htmlbox(pymupdf.Rect(255, 55 + 30 * i, 445, 78 + 30 * i), b)
    destino = tmp_path / "tablas.xlsx"
    assert pdf_a_excel(doc.tobytes(), str(destino)) == 6
    hoja = load_workbook(destino)["Página 1"]
    filas = [tuple(v for v in f if v is not None) for f in hoja.iter_rows(values_only=True)]
    assert [f for f in filas if f] == [("Concepto", "Importe"), ("Base imponible", 824.1), ("Total a pagar", 1028.26)]


def test_hoja_como_ilovepdf_casillas_fondo_imagen_y_codigo_de_barras(tmp_path):
    """Cada dato en su casilla, a su altura; el fondo gris pintado; las imágenes encima;
    y los caracteres raros de un código de barras no rompen el .xlsx."""
    from openpyxl import load_workbook

    from pdfcasero.conversion import pdf_a_excel

    doc = pymupdf.open()
    p = doc.new_page()
    foto = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    foto.set_rect(foto.irect, (200, 30, 30))
    p.insert_image(pymupdf.Rect(400, 40, 520, 80), pixmap=foto)
    p.insert_text((50, 120), "Nº de cuenta:", fontsize=9)
    p.insert_text((160, 120), "928912945", fontsize=9)
    p.insert_text((50, 134), "Fecha:", fontsize=9)
    p.insert_text((160, 134), "01/09/2025", fontsize=9)
    p.draw_rect(pymupdf.Rect(45, 160, 300, 180), color=None, fill=(0.7, 0.7, 0.7))
    p.insert_text((50, 173), "Total a Pagar", fontsize=9)
    p.insert_text((290 - pymupdf.get_text_length("1.028,26", fontsize=9), 173), "1.028,26", fontsize=9)
    p.insert_text((50, 220), "\x0c\x1cFIN", fontsize=9)
    destino = tmp_path / "factura.xlsx"
    pdf_a_excel(doc.tobytes(), str(destino))
    hoja = load_workbook(destino)["Página 1"]
    casillas = {c.value: c for f in hoja.iter_rows() for c in f if c.value is not None}
    cuenta, numero = casillas["Nº de cuenta:"], casillas["928912945"]
    assert cuenta.row == numero.row and numero.column > cuenta.column
    assert casillas["Fecha:"].row == cuenta.row + 1
    total = casillas[1028.26]
    assert total.row == casillas["Total a Pagar"].row and total.alignment.horizontal == "right"
    assert total.fill.fgColor.rgb.endswith("B2B2B2") and casillas["Total a Pagar"].fill.fgColor.rgb.endswith("B2B2B2")
    assert "FIN" in casillas
    import zipfile

    medios = [n for n in zipfile.ZipFile(destino).namelist() if n.startswith("xl/media/")]
    assert medios == ["xl/media/image1.png"]  # la imagen, encima de las casillas


def test_preparar_para_word_texto_girado_como_imagen_y_fuentes(tmp_path):
    import zipfile

    from pdfcasero.conversion import _clave_fuente, pdf_a_word, preparar_para_word

    assert _clave_fuente("ABCDEF+VodafoneRg-Bold") == _clave_fuente("Vodafone Rg") == "vodafonerg"
    doc = pymupdf.open()
    pagina = doc.new_page()
    pagina.insert_text((72, 100), "Factura de prueba con texto normal", fontname="cour", fontsize=11)
    pagina.insert_text((30, 700), "Texto vertical del margen", fontsize=8, rotate=90)
    datos, fuentes = preparar_para_word(doc.tobytes())
    preparada = pymupdf.open(stream=datos, filetype="pdf")[0]
    texto = preparada.get_text()
    assert "Factura de prueba" in texto and "vertical" not in texto  # el girado ahora es imagen
    assert len(preparada.get_images()) == 1
    # Courier es poco común en los PDF: se cambia por Courier New con su tamaño
    assert fuentes["courier"][0] == "Courier New" and 0.6 <= fuentes["courier"][1] <= 1.0

    origen = tmp_path / "prueba.pdf"
    doc.save(origen)
    destino = tmp_path / "prueba.docx"
    avances = []
    pdf_a_word(str(origen), str(destino), paginas=[0], aviso=lambda *a: avances.append(a))
    xml = zipfile.ZipFile(destino).read("word/document.xml").decode()
    assert "Factura de prueba" in xml and 'w:ascii="Courier New"' in xml
    assert 'w:lineRule="exact"' not in xml
    assert (1, 1, 1) in avances and (2, 1, 1) in avances
