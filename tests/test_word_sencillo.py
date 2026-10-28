"""El conversor «que fluye» para Google Docs: factura de ejemplo con datos, dirección y totales."""

import docx
import pymupdf

from pdfcasero.word_sencillo import pdf_a_word_sencillo


def factura() -> bytes:
    doc = pymupdf.open()
    p = doc.new_page()
    # Datos a la izquierda (etiqueta y valor) y dirección a la derecha
    for k, (etiqueta, valor) in enumerate([("Nº de factura:", "CI0925"), ("Fecha de emisión:", "01/09/2025")]):
        y = 100 + 14 * k
        p.insert_text((50, y), etiqueta, fontsize=8)
        p.insert_text((150, y), valor, fontsize=8)
    p.insert_text((330, 106), "LAURA BARCELO", fontsize=8)
    p.insert_text((330, 120), "03600 ELDA", fontsize=8)
    # Caja gris con totales alineados a la derecha
    p.draw_rect(pymupdf.Rect(43, 150, 265, 190), color=None, fill=(0.7, 0.7, 0.7))
    for k, (texto, importe) in enumerate([("Base imponible", "824,10"), ("Total a Pagar", "1.028,26")]):
        y = 165 + 15 * k
        p.insert_text((50, y), texto, fontsize=8)
        p.insert_text((258 - pymupdf.get_text_length(importe, fontsize=8), y), importe, fontsize=8)
    # Texto girado del margen (datos legales)
    p.insert_text((20, 600), "Vodafone España S.A.U. Madrid", fontsize=5, rotate=90)
    return doc.tobytes()


def test_factura_fluye_como_en_google_docs(tmp_path):
    destino = tmp_path / "factura.docx"
    avances = []
    pdf_a_word_sencillo(factura(), str(destino), avance=lambda i, n: avances.append((i, n)))
    assert avances == [(1, 1)]
    d = docx.Document(str(destino))

    def filas(tablas):
        for t in tablas:
            for r in t.rows:
                yield [c.text for c in r.cells], r
            for c in t._cells:
                yield from filas(c.tables)

    todas = list(filas(d.tables))
    # Google Docs pierde el texto de las tablas metidas dentro de otra tabla: no hay ninguna
    assert not any(c.tables for t in d.tables for c in t._cells)
    # ...y junta dos tablas seguidas en una: siempre hay un párrafo entre medias
    hijos = [e.tag.rsplit("}", 1)[1] for e in d.element.body]
    assert all(not (a == b == "tbl") for a, b in zip(hijos, hijos[1:]))
    # Datos a la izquierda y dirección a la derecha, cada uno en su columna; dentro de la
    # columna, etiqueta y valor en el mismo renglón
    assert any("Fecha de emisión:\t01/09/2025" in t[0] and "LAURA BARCELO" in t[1] for t, _ in todas)

    # El total en su caja gris (celdas sombreadas) y el importe alineado a la derecha
    texto, fila = next((t, r) for t, r in todas if "Total a Pagar" in t)
    assert 'w:fill="B2B2B2"' in fila._tr.xml
    importe = next(c for c in fila.cells if c.text == "1.028,26")
    assert importe.paragraphs[0].alignment == 2

    # El texto girado del margen, en horizontal en el pie de página
    assert "Vodafone España" in d.sections[0].footer.paragraphs[0].text


def test_solo_las_paginas_elegidas(tmp_path):
    doc = pymupdf.open()
    for n in range(3):
        doc.new_page().insert_text((72, 80), f"Página número {n + 1}", fontsize=12)
    destino = tmp_path / "una.docx"
    pdf_a_word_sencillo(doc.tobytes(), str(destino), paginas=[1])
    texto = "\n".join(p.text for p in docx.Document(str(destino)).paragraphs)
    assert "Página número 2" in texto and "número 1" not in texto and "número 3" not in texto
